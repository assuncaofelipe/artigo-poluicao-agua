"""Create a single-slide PowerPoint from a rendered poster image."""

from pathlib import Path
import sys

from pptx import Presentation
from pptx.util import Inches


def main(pdf_image: str, output: str) -> None:
    source = Path(pdf_image)
    destination = Path(output)
    destination.parent.mkdir(parents=True, exist_ok=True)

    presentation = Presentation()
    # The poster's physical size is 90 x 120 cm (35.433 x 47.244 in).
    presentation.slide_width = Inches(35.433071)
    presentation.slide_height = Inches(47.244094)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(
        str(source), 0, 0, width=presentation.slide_width, height=presentation.slide_height
    )
    presentation.save(destination)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("usage: build_presentation.py INPUT.png OUTPUT.pptx")
    main(sys.argv[1], sys.argv[2])
